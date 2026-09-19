import os

from langchain_core.documents import Document
from pypdf import PdfReader
from docx import Document as DocxDocument

from src.config import DATA_PATH


def load_documents():

    documents = []

    print(f"\nLooking for documents in: {DATA_PATH}")

    if not os.path.exists(DATA_PATH):

        print("ERROR: College documents folder does not exist!")

        return documents

    files = os.listdir(DATA_PATH)

    print(f"Files found: {files}")

    for filename in files:

        filepath = os.path.join(DATA_PATH, filename)

        # Skip folders
        if not os.path.isfile(filepath):
            continue

        try:

            # ---------- PDF FILES ----------

            if filename.lower().endswith(".pdf"):

                print(f"Loading PDF: {filename}")

                reader = PdfReader(filepath)

                for page_number, page in enumerate(reader.pages):

                    text = page.extract_text()

                    if text and text.strip():

                        documents.append(
                            Document(
                                page_content=text.strip(),
                                metadata={
                                    "source": filename,
                                    "page": page_number
                                }
                            )
                        )

            # ---------- TXT FILES ----------

            elif filename.lower().endswith(".txt"):

                print(f"Loading TXT: {filename}")

                with open(
                    filepath,
                    "r",
                    encoding="utf-8",
                    errors="ignore"
                ) as file:

                    text = file.read()

                if text and text.strip():

                    documents.append(
                        Document(
                            page_content=text.strip(),
                            metadata={
                                "source": filename,
                                "page": 0
                            }
                        )
                    )

            # ---------- DOCX FILES ----------

            elif filename.lower().endswith(".docx"):

                print(f"Loading DOCX: {filename}")

                doc = DocxDocument(filepath)

                text = "\n".join(
                    paragraph.text
                    for paragraph in doc.paragraphs
                    if paragraph.text.strip()
                )

                if text.strip():

                    documents.append(
                        Document(
                            page_content=text.strip(),
                            metadata={
                                "source": filename,
                                "page": 0
                            }
                        )
                    )

        except Exception as e:

            print(f"Could not load {filename}: {e}")

    print(f"\nTotal documents loaded: {len(documents)}")

    return documents